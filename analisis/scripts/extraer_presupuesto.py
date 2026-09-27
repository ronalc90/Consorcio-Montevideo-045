"""
Extrae el presupuesto 01-09-2026 (75MM, 8 meses) a datos estructurados.

Uso (desde la raíz del repo):
    python analisis/scripts/extraer_presupuesto.py

Genera:
    analisis/datos/presupuesto_2026_09.json   (hoja principal, ítem x CIV)
    analisis/datos/hojas/*.csv                (todas las hojas, valores; col 0 = A, fila n = fila n de Excel)

Requiere: pip install openpyxl
"""
import csv
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "fuentes" / "4. PRESUPUESTO 01-09-2026 75MM (8 MESES).xlsx"
OUT = ROOT / "analisis" / "datos"
HOJA = "PRESUPUESTO TODOS LOS CIV 84 NP"

CH_NAMES = {
    "1": "1. PRELIMINARES", "2": "2. PAVIMENTOS", "3": "3. ESPACIO PÚBLICO",
    "4": "4. SEÑALIZACIÓN Y DEMARCACIÓN", "5": "5. REDES HIDROSANITARIAS",
    "6": "6. REDES SECAS", "7": "7. DESVÍOS",
}


def num(v):
    if v in (None, ""):
        return None
    if isinstance(v, bool) or isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return v


def fmt_item(v):
    if isinstance(v, float):
        return f"{v:.3f}"
    return None if v in (None, "") else str(v).strip()


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "hojas").mkdir(exist_ok=True)
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)

    # 1) Todas las hojas a CSV
    for ws in wb.worksheets:
        fn = OUT / "hojas" / (ws.title.strip().replace(" ", "_").replace("/", "-") + ".csv")
        with open(fn, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            for r in ws.iter_rows(values_only=True):
                w.writerow(["" if v is None else v for v in r])

    # 2) Hoja principal estructurada
    rows = list(wb[HOJA].iter_rows(values_only=True))
    hdr_ids, hdr_seg = rows[2], rows[3]
    civ_cols = []  # (id, segmento, subgrupo, col_cant, col_val)  índices base 0
    for c in range(18, 75):
        cid = hdr_ids[c]
        if cid in (None, "") or "TOTAL" in str(cid).upper():
            continue
        civ_cols.append((str(cid), str(hdr_seg[c]), "2" if c < 32 else "5", c, c + 1))

    items, chapter, sub = [], None, None
    for n in range(6, 688):  # filas Excel 7..688
        r = rows[n]
        excel_row = n + 1
        code, ip, desc, und = r[1], r[2], r[5], r[6]
        desc_s = str(desc).strip() if desc not in (None, "") else ""
        if und in (None, "") and desc_s:
            if isinstance(ip, (int, float)) and float(ip).is_integer() and str(int(ip)) in CH_NAMES:
                chapter, sub = CH_NAMES[str(int(ip))], None
            elif str(ip).strip() in CH_NAMES:
                chapter, sub = CH_NAMES[str(ip).strip()], None
            elif excel_row >= 680 and desc_s == "ACEROS":
                chapter, sub = "8. ACTIVIDADES ACERO", None
            else:
                sub = desc_s
            continue
        if und in (None, ""):
            continue
        it = {
            "row": excel_row, "chapter": chapter, "subchapter": sub,
            "codigo_idu": None if code in (None, "") else str(code).strip(),
            "item_pago": fmt_item(ip), "esp_general": r[3], "esp_particular": r[4],
            "descripcion": desc_s, "und": str(und).strip(),
            "cant_contractual": num(r[7]), "cant_actualizada": num(r[8]), "delta_cant": num(r[9]),
            "vu_cd": num(r[10]), "vu_cd_aiu": num(r[11]),
            "valor_actualizado_aiu": num(r[12]), "valor_inicial_aiu": num(r[13]),
            "balance_may_men": num(r[14]), "incorporacion_np": num(r[15]), "adicion_total": num(r[16]),
            "total_sg2_cant": num(r[32]), "total_sg2_val": num(r[33]),
            "total_sg5_cant": num(r[75]), "total_sg5_val": num(r[76]),
            "total_cant": num(r[78]), "total_val": num(r[79]), "check": r[80], "incidencia": num(r[82]),
            "is_np": bool(re.match(r"^NP", str(ip or "").strip(), re.I)) or str(code).strip().upper() == "NO",
            "civ": {},
        }
        for cid, _seg, _sg, qc, vc in civ_cols:
            q, v = num(r[qc]), num(r[vc])
            if q not in (None, 0) or v not in (None, 0):
                it["civ"][cid] = {"cant": q, "valor": v}
        items.append(it)

    glob = {}
    for n in range(687, 708):
        r = rows[n]
        if r[1] not in (None, "") and r[12] not in (None, ""):
            glob[str(r[1]).strip()] = {"excel_row": n + 1, "valor": num(r[12]),
                                       "sg2": num(r[33]), "sg5": num(r[76]), "total_cols": num(r[79])}

    out = {
        "meta": {"fuente": XLSX.name, "hoja": HOJA, "aiu_factor": 0.31849, "plazo_meses_adicion": 8,
                 "nota": "Valores con AIU. Capítulo '8. ACTIVIDADES ACERO' (filas 682-687) NO suma en obras 58.196.933.800."},
        "civs": [{"id": c[0], "codigo_seg": c[1], "subgrupo": c[2]} for c in civ_cols],
        "items": items,
        "globales_hoja": glob,
    }
    with open(OUT / "presupuesto_2026_09.json", "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1, default=str)

    obras = sum((i["valor_actualizado_aiu"] or 0) for i in items if i["chapter"] != "8. ACTIVIDADES ACERO")
    print(f"{len(civ_cols)} CIVs | {len(items)} renglones | obras con AIU = {obras:,.0f} (debe ser 58,196,933,800)")


if __name__ == "__main__":
    main()
